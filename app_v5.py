# -*- coding: utf-8 -*-
import asyncio, json, os, random, re, subprocess, sys, threading, time
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urlparse, parse_qs
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import Workbook
from playwright.async_api import async_playwright

APP = "生意参谋赛道数据采集器 API V5"
ROOT = Path.home() / ".sycm_sector_miner_v5"
OUT = ROOT / "exports"
DBG = ROOT / "debug"
for p in (ROOT, OUT, DBG): p.mkdir(parents=True, exist_ok=True)
RANK = "/mc/mq/mkt/keyword/rank/pro.json"
RELATED = "/mc/mq/mkt/keyword/relate/analysis.json"

def compact(v): return re.sub(r"\s+", " ", str("" if v is None else v)).strip()
def metric(v): return v.get("value", "") if isinstance(v, dict) else ("" if v is None else v)
def comparison(v): return v.get("cycleCrc", "") if isinstance(v, dict) else ""

def upper(v):
    if isinstance(v, (int, float)): return float(v)
    s = compact(v).replace(",", "")
    def cv(x, u):
        n = float(x)
        return n * 10000 if u == "万" else n * 1000 if u == "千" else n
    m = re.search(r"([0-9.]+)\s*(万|千)?\s*[~～\-—到]\s*([0-9.]+)\s*(万|千)?", s)
    if m: return cv(m.group(3), m.group(4))
    m = re.search(r"([0-9.]+)\s*(万|千)?", s)
    return cv(m.group(1), m.group(2)) if m else None

def dates30(today=None):
    end = today or date.today()
    return [end - timedelta(days=i) for i in range(30)]

def day_range(d): return f"{d.isoformat()}|{d.isoformat()}"
def range30(today=None):
    ds = dates30(today)
    return f"{ds[-1].isoformat()}|{ds[0].isoformat()}"

def norm_url(s):
    s = compact(s)
    if s.startswith("https://") or s.startswith("http://"): return s
    if s.startswith("sycm.taobao.com/"): return "https://" + s
    if s.startswith("cm.taobao.com/"): return "https://sy" + s
    if s.startswith("/mc/free/search_rank"): return "https://sycm.taobao.com" + s
    if s.startswith("/search_rank"): return "https://sycm.taobao.com/mc/free" + s
    if s.startswith("search_rank"): return "https://sycm.taobao.com/mc/free/" + s
    return s

def ctx(url):
    q = parse_qs(urlparse(url).query)
    g = lambda k: (q.get(k) or [""])[0]
    return {"parentCateId": g("parentCateId"), "cateId": g("cateId"), "cateFlag": g("cateFlag")}

def rank_params(cc, d, page=1):
    return {"dateRange": day_range(d), "dateType": "day", "page": page, "pageSize": 50,
            "order": "desc", "marketVersion": "free", "parentCateId": cc["parentCateId"],
            "cateId": cc["cateId"], "cateFlag": cc["cateFlag"], "kwType": "search",
            "rankType": "hot", "orderBy": "seIpvUvHits", "device": 0}

def related_params(keyword, dr, page=1):
    return {"dateRange": dr, "dateType": "day", "pageSize": 100, "page": page,
            "order": "desc", "orderBy": "seIpvUvHits", "keyWord": keyword,
            "cycleFlag": "yearSync", "rankType": "related", "marketVersion": "free"}

class Store:
    def __init__(self, path):
        self.path = Path(path); self.wb = Workbook()
        self.rank = self.wb.active; self.rank.title = "图1_30天原始"
        self.rank.append(["数据日期","页码","当日排名","搜索词","搜索人气","点击率","支付转化率","搜索增速"])
        self.summary = self.wb.create_sheet("主词汇总")
        self.summary.append(["搜索词","出现天数","30天最高搜索人气","30天最低搜索人气","最近出现日期","最近搜索人气"])
        self.related = self.wb.create_sheet("图2_关联词")
        self.related.append(["数据周期","对比方式","来源搜索词","关联搜索词","搜索人气","搜索人气同比","点击率","点击率同比","支付转化率","支付转化率同比","支付买家数","支付买家数同比","需求供给比","需求供给比同比","天猫商品点击占比","天猫商品点击占比同比"])
        for ws in (self.rank, self.summary, self.related): ws.freeze_panes = "A2"
        self.wb.save(self.path)
    def add_rank(self, rows):
        for r in rows: self.rank.append(r)
        self.wb.save(self.path)
    def add_summary(self, rows):
        for r in rows: self.summary.append(r)
        self.wb.save(self.path)
    def add_related(self, rows):
        for r in rows: self.related.append(r)
        self.wb.save(self.path)

class Collector:
    def __init__(self, url, main_min, related_min, output_dir, log_cb, status_cb):
        self.url = norm_url(url); self.main_min = int(main_min); self.related_min = int(related_min)
        self.output_dir = Path(output_dir); self.log_cb = log_cb; self.status_cb = status_cb
        self.stop_flag = False; self.days = dates30(); self.related_range = range30()
    def log(self, s): self.log_cb(f"[{time.strftime('%H:%M:%S')}] {s}")

    async def api(self, page, endpoint, params, referer):
        js = """async a=>{const u=new URL(a.e,location.origin);Object.entries(a.p).forEach(([k,v])=>u.searchParams.set(k,String(v)));try{const r=await fetch(u,{credentials:'include',headers:{accept:'*/*','sycm-referer':a.ref}});const t=await r.text();let x=null;try{x=JSON.parse(t)}catch(e){};return{status:r.status,json:x,text:x?null:t.slice(0,1000),url:u.toString()}}catch(e){return{status:0,json:null,text:String(e),url:u.toString()}}}"""
        return await page.evaluate(js, {"e": endpoint, "p": params, "ref": referer})

    async def api_retry(self, page, endpoint, params, referer):
        last = None
        for i in range(3):
            if self.stop_flag: raise RuntimeError("已停止")
            last = await self.api(page, endpoint, params, referer)
            x = last.get("json") if isinstance(last, dict) else None
            if isinstance(x, dict) and int(x.get("code", -1)) == 0:
                d = x.get("data") or {}
                return list(d.get("data") or []), int(d.get("recordCount") or 0)
            if i < 2:
                self.log(f"接口失败，第{i+1}次重试"); await asyncio.sleep(1.3)
        f = DBG / f"api_error_{int(time.time())}.json"
        f.write_text(json.dumps(last, ensure_ascii=False, indent=2), encoding="utf-8")
        msg = ""
        if isinstance(last, dict) and isinstance(last.get("json"), dict): msg = str(last["json"].get("message") or "")
        raise RuntimeError(f"生意参谋接口返回失败{('：'+msg) if msg else ''}。调试文件：{f}")

    async def ensure_analysis_page(self, page, kw):
        dr = self.related_range
        url = f"https://sycm.taobao.com/mc/free/search_analysis?keyWord={kw}&dateType=day&dateRange={dr}&searchAnalysisRadio=yearSync"
        await page.goto(url, wait_until="domcontentloaded", timeout=30000)
        await asyncio.sleep(1.5)
        if "/mc/free/search_analysis" not in page.url:
            raise RuntimeError("未进入搜索分析页")

    async def preflight(self, page, cc):
        d = self.days[0]
        self.log(f"自检1/2：图1 {d.isoformat()}")
        rows, _ = await self.api_retry(page, RANK, rank_params(cc, d, 1), "/mc/free/search_rank")
        first = next((x for x in rows if isinstance(x, dict) and compact(metric(x.get("searchWord")))), None)
        if not first: raise RuntimeError("图1无有效搜索词")
        kw = compact(metric(first.get("searchWord")))
        self.log(f"图1通过：{kw}")
        self.log("自检2/2：图2 30天/年同比")
        await self.ensure_analysis_page(page, kw)
        rows2, _ = await self.api_retry(page, RELATED, related_params(kw, self.related_range, 1), "/mc/free/search_analysis")
        if rows2 and not isinstance(rows2[0], dict): raise RuntimeError("图2返回结构异常")
        self.log(f"图2通过：返回 {len(rows2)} 条")
        await page.goto(self.url, wait_until="domcontentloaded", timeout=30000); await asyncio.sleep(.7)

    async def collect_rank_day(self, page, cc, d, store, candidates):
        page_no = 1; day_rank = 0; kept = 0
        while not self.stop_flag and page_no <= 100:
            rows, total = await self.api_retry(page, RANK, rank_params(cc, d, page_no), "/mc/free/search_rank")
            if not rows: break
            out = []; hit = False
            for item in rows:
                if not isinstance(item, dict): continue
                kw = compact(metric(item.get("searchWord")))
                if not kw: continue
                day_rank += 1
                pop = metric(item.get("seIpvUvHits")); u = upper(pop)
                if u is not None and u < self.main_min: hit = True; break
                out.append([d.isoformat(), page_no, day_rank, kw, pop, metric(item.get("freeClkRate")), metric(item.get("payRate")), metric(item.get("seUvEx"))])
                kept += 1
                rec = candidates.setdefault(kw, {"days": set(), "values": [], "latest_date": "", "latest_pop": ""})
                rec["days"].add(d.isoformat())
                if u is not None: rec["values"].append(u)
                if not rec["latest_date"] or d.isoformat() > rec["latest_date"]:
                    rec["latest_date"] = d.isoformat(); rec["latest_pop"] = pop
            if out: store.add_rank(out)
            if hit or len(rows) < 50 or (total and page_no * 50 >= total): break
            page_no += 1; await asyncio.sleep(random.uniform(.25,.5))
        self.log(f"图1 {d.isoformat()}：{kept} 个词")

    async def collect_related(self, page, kw):
        await self.ensure_analysis_page(page, kw)
        result = []; page_no = 1
        while not self.stop_flag and page_no <= 100:
            p = related_params(kw, self.related_range, page_no)
            rows, total = await self.api_retry(page, RELATED, p, "/mc/free/search_analysis")
            if not rows: break
            hit = False
            for item in rows:
                if not isinstance(item, dict): continue
                pop = metric(item.get("seIpvUvHits")); u = upper(pop)
                if u is not None and u < self.related_min: hit = True; break
                result.append([self.related_range, "年同比", kw, metric(item.get("relatedSekeyword")), pop,
                               comparison(item.get("seIpvUvHits")), metric(item.get("freeClkRate")), comparison(item.get("freeClkRate")),
                               metric(item.get("payConvRate")), comparison(item.get("payConvRate")), metric(item.get("payByrCnt")), comparison(item.get("payByrCnt")),
                               metric(item.get("simWeight")), comparison(item.get("simWeight")), metric(item.get("tmaoClickRatio")), comparison(item.get("tmaoClickRatio"))])
            if hit or len(rows) < 100 or (total and page_no * 100 >= total): break
            page_no += 1; await asyncio.sleep(random.uniform(.3,.6))
        return result

    async def run(self):
        self.output_dir.mkdir(parents=True, exist_ok=True)
        path = self.output_dir / f"生意参谋_API_V5_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"; store = Store(path)
        async with async_playwright() as p:
            browser = await p.chromium.connect_over_cdp("http://127.0.0.1:9222")
            if not browser.contexts: raise RuntimeError("没有连接到采集专用 Chrome")
            context = browser.contexts[0]; pages = context.pages
            page = next((x for x in pages if "sycm.taobao.com" in x.url), pages[-1] if pages else await context.new_page())
            await page.goto(self.url, wait_until="domcontentloaded", timeout=30000); await asyncio.sleep(.8)
            if "sycm.taobao.com" not in page.url: raise RuntimeError("请先登录生意参谋")
            cc = ctx(self.url or page.url)
            if not cc["cateId"]: cc = ctx(page.url)
            if not cc["cateId"]: raise RuntimeError("没有识别到 cateId")
            self.log(f"类目 cateId={cc['cateId']}｜图1逐日30天｜图2={self.related_range}｜年同比")
            await self.preflight(page, cc)
            candidates = {}
            for i, d in enumerate(self.days, 1):
                if self.stop_flag: break
                self.status_cb(f"图1：第{i}/30天｜{d.isoformat()}")
                await self.collect_rank_day(page, cc, d, store, candidates)
                await asyncio.sleep(random.uniform(.3,.6))
            summary = []
            for kw, rec in sorted(candidates.items(), key=lambda kv:max(kv[1]["values"] or [0]), reverse=True):
                vals = rec["values"]; summary.append([kw,len(rec["days"]),max(vals) if vals else "",min(vals) if vals else "",rec["latest_date"],rec["latest_pop"]])
            if summary: store.add_summary(summary)
            self.log(f"图1完成：去重主词 {len(candidates)} 个")
            for i, kw in enumerate(candidates.keys(), 1):
                if self.stop_flag: break
                self.status_cb(f"图2：{i}/{len(candidates)}｜{kw}")
                rows = await self.collect_related(page, kw)
                if rows: store.add_related(rows)
                self.log(f"图2 {kw}：{len(rows)} 条")
                await asyncio.sleep(random.uniform(.4,.8))
        self.log(f"完成：{path}"); return str(path)

class App:
    def __init__(self, root):
        self.root=root; root.title(APP); root.geometry("1020x740"); self.worker=None; self.collector=None
        f=ttk.Frame(root,padding=16); f.pack(fill="both",expand=True)
        ttk.Label(f,text=APP,font=("Microsoft YaHei UI",18,"bold")).grid(row=0,column=0,columnspan=4,sticky="w",pady=(0,14))
        ttk.Label(f,text="搜索排行 URL").grid(row=1,column=0,sticky="w"); self.url=tk.StringVar(); ttk.Entry(f,textvariable=self.url,width=92).grid(row=1,column=1,columnspan=3,sticky="ew",pady=5)
        ttk.Label(f,text="图1主词最低搜索人气").grid(row=2,column=0,sticky="w"); self.mn=tk.IntVar(value=300); ttk.Entry(f,textvariable=self.mn,width=12).grid(row=2,column=1,sticky="w")
        ttk.Label(f,text="图2关联词最低搜索人气").grid(row=2,column=2,sticky="w"); self.rn=tk.IntVar(value=100); ttk.Entry(f,textvariable=self.rn,width=12).grid(row=2,column=3,sticky="w")
        ttk.Label(f,text="图1周期").grid(row=3,column=0,sticky="w"); ttk.Label(f,text="过去30天（逐日抓取）").grid(row=3,column=1,sticky="w")
        ttk.Label(f,text="图2周期 / 对比").grid(row=3,column=2,sticky="w"); ttk.Label(f,text="过去30天 / 年同比").grid(row=3,column=3,sticky="w")
        ttk.Label(f,text="导出目录").grid(row=4,column=0,sticky="w"); self.od=tk.StringVar(value=str(OUT)); ttk.Entry(f,textvariable=self.od).grid(row=4,column=1,columnspan=2,sticky="ew",pady=5); ttk.Button(f,text="选择目录",command=self.pick).grid(row=4,column=3,sticky="w")
        b=ttk.Frame(f); b.grid(row=5,column=0,columnspan=4,sticky="w",pady=(12,8))
        for t,cmd in [("启动采集专用 Chrome",self.chrome),("测试 Chrome",self.test),("开始采集",self.start),("停止",self.stop),("打开导出目录",self.openout)]: ttk.Button(b,text=t,command=cmd).pack(side="left",padx=4)
        self.st=tk.StringVar(value="V5：图2直接使用已验证参数，不再依赖请求捕获"); ttk.Label(f,textvariable=self.st).grid(row=6,column=0,columnspan=4,sticky="w")
        ttk.Label(f,text="运行日志").grid(row=7,column=0,columnspan=4,sticky="w",pady=(10,2)); self.box=tk.Text(f,height=25,wrap="word"); self.box.grid(row=8,column=0,columnspan=4,sticky="nsew")
        ttk.Label(f,text="首次：启动专用 Chrome → 登录生意参谋 → 粘贴图1完整网址 → 开始采集。",foreground="#555").grid(row=9,column=0,columnspan=4,sticky="w",pady=(10,0)); f.rowconfigure(8,weight=1)
        for i in range(4): f.columnconfigure(i,weight=1)
    def log(self,s): self.root.after(0,lambda x=s:(self.box.insert("end",x+"\n"),self.box.see("end")))
    def status(self,s): self.root.after(0,lambda x=s:self.st.set(x))
    def pick(self):
        p=filedialog.askdirectory(initialdir=self.od.get())
        if p:self.od.set(p)
    def openout(self): Path(self.od.get()).mkdir(parents=True,exist_ok=True); os.startfile(self.od.get())
    def chrome(self):
        cs=[Path(os.environ.get("ProgramFiles",""))/"Google/Chrome/Application/chrome.exe",Path(os.environ.get("ProgramFiles(x86)",""))/"Google/Chrome/Application/chrome.exe",Path(os.environ.get("LocalAppData",""))/"Google/Chrome/Application/chrome.exe"]
        c=next((x for x in cs if x.exists()),None)
        if not c: messagebox.showerror("未找到 Chrome","请先安装 Google Chrome"); return
        prof=ROOT/"chrome-profile"; prof.mkdir(parents=True,exist_ok=True)
        subprocess.Popen([str(c),"--remote-debugging-port=9222","--remote-allow-origins=*",f"--user-data-dir={prof}","https://sycm.taobao.com/"])
        self.st.set("已启动专用 Chrome；首次请手动登录")
    def test(self):
        def w():
            async def go():
                async with async_playwright() as p:
                    b=await p.chromium.connect_over_cdp("http://127.0.0.1:9222"); return [pg.url for c in b.contexts for pg in c.pages]
            try:
                u=asyncio.run(go()); self.log("Chrome 连接成功："+" | ".join(u[:3])); self.status("Chrome 连接成功"); self.root.after(0,lambda:messagebox.showinfo("成功","Chrome 已连接"))
            except Exception as e:
                er=str(e); self.log("Chrome 连接失败："+er); self.root.after(0,lambda x=er:messagebox.showerror("失败","请先启动采集专用 Chrome。\n\n"+x))
        threading.Thread(target=w,daemon=True).start()
    def start(self):
        if self.worker and self.worker.is_alive(): messagebox.showwarning("提示","任务正在运行"); return
        if not compact(self.url.get()): messagebox.showwarning("缺少 URL","请粘贴图1完整网址"); return
        self.collector=Collector(self.url.get(),self.mn.get(),self.rn.get(),self.od.get(),self.log,self.status); self.worker=threading.Thread(target=self._run,daemon=True); self.worker.start()
    def _run(self):
        try:
            p=asyncio.run(self.collector.run()); self.status("完成："+p); self.root.after(0,lambda x=p:messagebox.showinfo("完成","文件已保存：\n"+x))
        except Exception as e:
            er=str(e); self.log("运行失败："+er); self.status("运行失败"); self.root.after(0,lambda x=er:messagebox.showerror("运行失败",x))
    def stop(self):
        if self.collector:self.collector.stop_flag=True; self.status("正在停止…")

def selftest():
    assert dates30(date(2026,9,15))[0].isoformat()=="2026-09-15"
    assert dates30(date(2026,9,15))[-1].isoformat()=="2026-08-17"
    assert range30(date(2026,9,15))=="2026-08-17|2026-09-15"
    p=related_params("剃须刀",range30(date(2026,9,15)),1)
    assert p["cycleFlag"]=="yearSync" and p["rankType"]=="related" and p["pageSize"]==100
    assert norm_url("cm.taobao.com/mc/free/search_rank?cateId=1").startswith("https://sycm.taobao.com/")
    assert upper("120万 ~ 250万")==2500000
    print("SELF_TEST_OK")

if __name__=="__main__":
    if "--self-test" in sys.argv:selftest()
    else:
        r=tk.Tk(); App(r); r.mainloop()
